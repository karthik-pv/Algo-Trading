"""Convert setup occurrence CSVs to Excel with colored Trend cells.

Colors: UT = dark green, RT (retracement up) = light green,
        DT = dark red,   RD (retracement down) = light red.
Mixed/neutral trend values get no fill.

Run:  python -m forecasting.setup_xlsx
Reads data/insight/csv/setup_occurrences*.csv, writes .xlsx alongside.
"""

import csv
import glob
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_DIR = os.path.join(_REPO_ROOT, "data", "insight", "csv")

FILLS = {
    "UT": PatternFill("solid", start_color="548235"),   # dark green
    "RT": PatternFill("solid", start_color="C6EFCE"),   # light green
    "DT": PatternFill("solid", start_color="C00000"),   # dark red
    "RD": PatternFill("solid", start_color="FFC7CE"),   # light red
}
FONTS = {
    "UT": Font(color="FFFFFF", bold=True),
    "RT": Font(color="006100"),
    "DT": Font(color="FFFFFF", bold=True),
    "RD": Font(color="9C0006"),
}


def trend_key(value):
    v = (value or "").strip()
    if v.startswith("RT_UP") or v.startswith("RTU"):
        return "RT"
    if v.startswith("RT_DOWN") or v.startswith("RTD"):
        return "RD"
    if v.startswith("RT"):
        return "RT"
    if v.startswith("RD"):
        return "RD"
    if v.startswith("UT"):
        return "UT"
    if v.startswith("DT"):
        return "DT"
    return None


def convert(csv_path):
    xlsx_path = os.path.splitext(csv_path)[0] + ".xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = os.path.splitext(os.path.basename(csv_path))[0][:31]
    with open(csv_path, newline="") as f:
        for i, row in enumerate(csv.reader(f)):
            ws.append(row)
            if i == 0:
                continue
            key = trend_key(row[3] if len(row) > 3 else "")
            if key:
                cell = ws.cell(row=i + 1, column=4)
                cell.fill = FILLS[key]
                cell.font = FONTS[key]
    # sensible column widths
    for col in ws.columns:
        width = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    return xlsx_path, ws.max_row - 1


def main():
    for path in sorted(glob.glob(os.path.join(CSV_DIR, "setup_occurrences*.csv"))):
        if path.endswith(".xlsx"):
            continue
        out, n = convert(path)
        print("%s -> %s (%d rows)" % (os.path.basename(path), os.path.basename(out), n))


if __name__ == "__main__":
    main()
