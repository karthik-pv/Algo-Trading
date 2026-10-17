"""Live browser view of the setup trade-case table.

Serves data/insight/setups/setup_pa_tradecase.xlsx as an auto-refreshing HTML
table at http://localhost:8917 . Reads the xlsx fresh on every request, so any
update to the file shows up within the refresh interval.

Run:  .venv/Scripts/python.exe scripts/view_setup_table.py  [port]
Stop: close the terminal / kill the process.
"""

import html
import os
import sys
from http.server import BaseHTTPRequestHandler, HTTPServer

from openpyxl import load_workbook

XLSX = os.path.join("data", "insight", "setups", "setup_pa_tradecase.xlsx")
PORT = int(sys.argv[1]) if len(sys.argv) > 1 else 8917

PAGE = """<!DOCTYPE html>
<html><head><meta charset="utf-8">
<meta http-equiv="refresh" content="5">
<title>Setup / PA / Trade Case</title>
<style>
 body {{ font-family: 'Segoe UI', sans-serif; margin: 16px; background: #fafafa; }}
 h2 {{ margin: 0 0 4px 0; }}
 .meta {{ color: #666; font-size: 12px; margin-bottom: 10px; }}
 table {{ border-collapse: collapse; width: 100%; background: #fff; }}
 th, td {{ border: 1px solid #cfd8dc; padding: 8px 10px; vertical-align: top;
          text-align: left; }}
 th {{ background: #d9e1f2; position: sticky; top: 0; }}
 tr:nth-child(even) td {{ background: #f4f7fb; }}
 td:nth-child(5) {{ white-space: pre; font-family: Consolas, monospace; }}
 td:nth-child(6) {{ font-size: 12px; color: #333; }}
</style></head><body>
<h2>Setup / Price Action / Trade Case</h2>
<div class="meta">auto-refreshes every 5s &middot; file last modified: {mtime}</div>
<table>{rows}</table>
</body></html>"""


def render():
    try:
        wb = load_workbook(XLSX, data_only=True)
        ws = wb.active
        rows_html = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if all(v is None for v in row):
                continue
            tag = "th" if r_i == 1 else "td"
            cells = "".join("<%s>%s</%s>" % (tag, html.escape(str(v)) if v is not None else "", tag)
                            for v in row)
            rows_html.append("<tr>%s</tr>" % cells)
        mtime = datetime_str()
        return PAGE.format(rows="".join(rows_html), mtime=mtime)
    except Exception as e:
        return "<html><body><h2>Waiting for file... (%s)</h2></body></html>" % html.escape(str(e))


def datetime_str():
    try:
        import datetime as dt
        m = os.path.getmtime(XLSX)
        return dt.datetime.fromtimestamp(m).strftime("%d %b %Y %H:%M:%S")
    except OSError:
        return "n/a"


class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        body = render().encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *a):
        pass


if __name__ == "__main__":
    srv = HTTPServer(("127.0.0.1", PORT), Handler)
    print("Serving %s at http://localhost:%d  (Ctrl+C to stop)" % (XLSX, PORT))
    srv.serve_forever()
